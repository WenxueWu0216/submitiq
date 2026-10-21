"""Load a submission package: every file becomes a Document with pages of words and text."""
from __future__ import annotations

import hashlib
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pdfplumber
from openpyxl import load_workbook

from .ocr import ocr_page
from .textlayout import Line, Word, group_lines

MIN_TEXT_CHARS = 25   # below this a PDF page is treated as image-only and sent to OCR


@dataclass
class Page:
    number: int                         # 1-based
    words: list[Word]
    ocr: bool = False
    ocr_conf: float = 1.0
    info: dict = field(default_factory=dict)
    _lines: list[Line] | None = None

    @property
    def lines(self) -> list[Line]:
        if self._lines is None:
            self._lines = group_lines(self.words)
        return self._lines

    @property
    def text(self) -> str:
        return "\n".join(l.text() for l in self.lines)


@dataclass
class Sheet:
    name: str
    rows: list[list[Any]]               # raw cell values, row 1 first


@dataclass
class Document:
    name: str
    path: Path
    ext: str
    sha256: str
    pages: list[Page] = field(default_factory=list)
    sheets: list[Sheet] = field(default_factory=list)
    raw_text: str | None = None          # email / plain text
    doc_type: str = "other"
    timings_ms: dict[str, float] = field(default_factory=dict)

    @property
    def is_scanned(self) -> bool:
        return any(p.ocr for p in self.pages)

    def text(self, max_chars: int | None = None) -> str:
        if self.raw_text is not None:
            out = self.raw_text
        elif self.sheets:
            parts = []
            for sh in self.sheets:
                parts.append(f"[sheet {sh.name}]")
                for i, row in enumerate(sh.rows, start=1):
                    if any(c not in (None, "") for c in row):
                        parts.append(f"R{i}: " + " | ".join("" if c is None else str(c) for c in row))
            out = "\n".join(parts)
        else:
            out = "\n".join(f"--- page {p.number}{' (OCR)' if p.ocr else ''} ---\n{p.text}" for p in self.pages)
        return out[:max_chars] if max_chars else out


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_pdf(path: Path, sha: str) -> Document:
    doc = Document(path.name, path, ".pdf", sha)
    t_text = t_ocr = 0.0
    with pdfplumber.open(str(path)) as pdf:
        for i, page in enumerate(pdf.pages):
            t0 = time.perf_counter()
            raw = page.extract_words(x_tolerance=1.5, keep_blank_chars=False)
            words = [Word(w["text"], float(w["x0"]), float(w["x1"]), float(w["top"]), float(w["bottom"])) for w in raw]
            t_text += (time.perf_counter() - t0) * 1000
            if sum(len(w.text) for w in words) >= MIN_TEXT_CHARS:
                doc.pages.append(Page(i + 1, words))
                continue
            ocr_words, conf, info = ocr_page(path, i, sha)
            t_ocr += info.get("ocr_ms", 0.0)
            doc.pages.append(Page(i + 1, ocr_words, ocr=True, ocr_conf=conf, info=info))
    doc.timings_ms = {"text": t_text, "ocr": t_ocr}
    return doc


def load_xlsx(path: Path, sha: str) -> Document:
    t0 = time.perf_counter()
    wb = load_workbook(str(path), data_only=True, read_only=True)
    sheets = [Sheet(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    wb.close()
    doc = Document(path.name, path, ".xlsx", sha, sheets=sheets)
    doc.timings_ms = {"text": (time.perf_counter() - t0) * 1000}
    return doc


def load_package(pkg: Path) -> list[Document]:
    docs: list[Document] = []
    for path in sorted(pkg.iterdir()):
        if path.name == "ground_truth.json" or path.name.startswith("."):
            continue
        sha = _sha(path)
        ext = path.suffix.lower()
        if ext == ".pdf":
            docs.append(load_pdf(path, sha))
        elif ext in (".xlsx", ".xlsm"):
            docs.append(load_xlsx(path, sha))
        elif ext in (".txt", ".eml"):
            docs.append(Document(path.name, path, ext, sha, raw_text=path.read_text(errors="replace")))
        else:
            docs.append(Document(path.name, path, ext, sha, raw_text=""))
    return docs
