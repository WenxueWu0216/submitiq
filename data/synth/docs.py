"""Render the documents of one submission package from the visible ground truth."""
from __future__ import annotations

import random
from datetime import timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from submitiq.schema import Submission

from . import formats as F
from .pdfdraw import Column, PdfDoc
from .truth import World


def _addr_line(street, city, state, zip_) -> str:
    return ", ".join(p for p in [street or "", f"{city or ''}, {state or ''} {zip_ or ''}".strip(", ")] if p)


# --------------------------------------------------------------------------- email

def render_email(w: World, gold: Submission, rng: random.Random) -> str:
    ins, br, pol = gold.insured, gold.broker, gold.policy
    eff = w.email_effective_date or pol.effective_date
    name = w.name_variants["email"]
    received = pol.submission_received_date
    needed = pol.quote_needed_by
    greeting = rng.choice(["Hi team,", "Hello,", "Good morning,", "Hi underwriting,"])
    lines_req = rng.choice(["commercial property", "property (building, BPP, BI)", "a property quote"])
    tiv_note = ""
    if w.mentions_tiv_in_email and gold.totals.total_tiv:
        approx = round(gold.totals.total_tiv / 1_000_000, 1)
        tiv_note = f" Total values are roughly ${approx}M across {len(gold.locations)} location(s)."
    renewal_note = rng.choice([
        "The incumbent is pushing a large rate increase and the insured wants options.",
        "This is a new venture for us; the account has been with the same agent for years.",
        "Loss runs are attached and are currently valued.",
        "Please let me know if you need anything else to quote.",
    ])
    needed_phrase = rng.choice([
        f"We need indications by {needed.strftime('%B %-d')} if at all possible.",
        f"Quote needed by {needed.strftime('%m/%d/%Y')}.",
        f"Could you get us terms no later than {needed.strftime('%A, %B %-d, %Y')}?",
    ]) if needed else ""
    eff_phrase = rng.choice([
        f"with a requested effective date of {eff.strftime('%m/%d/%Y')}",
        f"effective {eff.strftime('%B %-d, %Y')}",
        f"with an inception date of {eff.strftime('%-m/%-d/%y')}",
    ])
    body = f"""{greeting}

Please find attached a new submission for {name} for {lines_req}, {eff_phrase}.{tiv_note}
{ins.business_description}.

{renewal_note} {needed_phrase}

Attached: application, loss runs and statement of values.

Thanks,
{br.producer_name}
{br.agency_name}
{br.producer_phone}
{br.producer_email}
"""
    subject = rng.choice([f"New submission - {name}", f"Property submission: {name}",
                          f"RE: {name} - property quote request"])
    hdr = (f"From: {br.producer_name} <{br.producer_email}>\n"
           f"To: submissions@carrier-example.com\n"
           f"Date: {received.strftime('%a, %d %b %Y')} 09:{rng.randint(10, 59)}:00 -0500\n"
           f"Subject: {subject}\n")
    return hdr + "\n" + body


# --------------------------------------------------------------------------- application

def _visible(w: World, path: str, value):
    return None if path in w.omitted else value


def _app_values(w: World, g: Submission, rng: random.Random) -> dict[str, str]:
    m = w.messy_formats
    ins, pol, cov, br = g.insured, g.policy, g.coverage, g.broker
    return {
        "agency": br.agency_name or "", "producer": br.producer_name or "",
        "producer_phone": br.producer_phone or "", "producer_email": br.producer_email or "",
        "named_insured": ins.named_insured or "", "dba": ins.dba or "", "fein": ins.fein or "",
        "entity": F.ENTITY_LABEL.get(ins.entity_type or "", ""),
        "street": ins.mailing_address.street or "", "city": ins.mailing_address.city or "",
        "state": ins.mailing_address.state or "", "zip": ins.mailing_address.zip or "",
        "description": ins.business_description or "", "naics": ins.naics_code or "",
        "years": F.intval(ins.years_in_business), "revenue": F.money(ins.annual_revenue, rng, m),
        "employees": F.intval(ins.employee_count), "website": ins.website or "",
        "contact": ins.contact_name or "", "contact_phone": ins.contact_phone or "",
        "contact_email": ins.contact_email or "",
        "eff": F.date_str(pol.effective_date, rng, m), "exp": F.date_str(pol.expiration_date, rng, m),
        "transaction": (pol.transaction_type or "").title(), "prior_carrier": pol.prior_carrier or "",
        "bldg_limit": F.money(cov.building_limit, rng, m), "bpp_limit": F.money(cov.bpp_limit, rng, m),
        "bi_limit": F.money(cov.business_income_limit, rng, m), "deductible": F.money(cov.deductible, rng, False),
        "valuation": F.VALUATION_LABEL.get(cov.valuation or "", ""),
        "coins": f"{cov.coinsurance_pct}%" if cov.coinsurance_pct else "",
        "form": F.FORM_LABEL.get(cov.cause_of_loss_form or "", ""),
        "flood": F.yes_no(cov.flood_requested, rng, m), "eq": F.yes_no(cov.earthquake_requested, rng, m),
        "tiv": F.money(w.app_stated_tiv, rng, m), "nlocs": str(len(g.locations)),
    }


def _premises_rows(w: World, g: Submission, rng: random.Random) -> list[list[str]]:
    m = w.messy_formats
    rows = []
    for loc in g.locations[:3]:
        yb = w.app_year_built.get(loc.loc_number, loc.year_built)
        rows.append([str(loc.loc_number), loc.street or "", loc.city or "", loc.state or "", loc.zip or "",
                     F.construction(loc.construction, rng, m), F.intval(yb), F.intval(loc.stories),
                     F.intcomma(loc.square_feet), F.yes_no(loc.sprinklered, rng, m), F.intval(loc.protection_class),
                     F.ROOF_LABEL.get(loc.roof_type or "", ""), F.intval(loc.roof_year),
                     F.ALARM_LABEL.get(loc.alarm_type or "", "")])
    return rows


PREMISES_COLS = [Column("Loc #", 30), Column("Street Address", 130), Column("City", 80), Column("ST", 22),
                 Column("Zip", 38), Column("Construction", 95), Column("Yr Built", 40), Column("Stories", 36),
                 Column("Sq Ft", 46, "right"), Column("Sprk", 28), Column("PC", 22), Column("Roof Type", 72),
                 Column("Roof Yr", 38), Column("Alarm", 70)]


def _sign_page(d: PdfDoc, v: dict[str, str], g: Submission) -> None:
    d.heading("PRIOR CARRIER / LOSS HISTORY")
    d.text(40, "Loss history for the past five years: see attached carrier loss runs.")
    d.heading("SIGNATURE")
    d.text(40, "The undersigned authorized representative of the applicant represents that the statements")
    d.text(40, "set forth in this application are true. This form is a synthetic look-alike for testing.")
    d.y -= 10
    d.text(40, f"Applicant's Signature: /s/ {v['contact']}")
    sign_date = g.policy.submission_received_date - timedelta(days=3) if g.policy.submission_received_date else None
    d.text(40, f"Date: {sign_date.strftime('%m/%d/%Y') if sign_date else ''}")


def render_application_a(w: World, g: Submission, rng: random.Random) -> bytes:
    v = _app_values(w, g, rng)
    d = PdfDoc()
    d.text(40, "COMMERCIAL INSURANCE APPLICATION", size=14, bold=True)
    d.text(40, "Applicant Information Section / Property Section  (ACORD 125/140-style, synthetic)", size=8)
    d.heading("AGENCY")
    d.kv_line([("Agency", v["agency"])], [40])
    d.kv_line([("Producer", v["producer"]), ("Producer Phone", v["producer_phone"])], [40, 330])
    d.kv_line([("Producer Email", v["producer_email"])], [40])
    d.heading("APPLICANT INFORMATION")
    d.kv_line([("Named Insured", v["named_insured"])], [40])
    d.kv_line([("DBA", v["dba"])], [40])
    d.kv_line([("FEIN", v["fein"]), ("Entity Type", v["entity"])], [40, 330])
    d.kv_line([("Mailing Address", _addr_line(v["street"], v["city"], v["state"], v["zip"]))], [40])
    d.kv_line([("Nature of Business", v["description"])], [40])
    d.kv_line([("NAICS Code", v["naics"]), ("Years in Business", v["years"])], [40, 330])
    d.kv_line([("Annual Revenue", v["revenue"]), ("Number of Employees", v["employees"])], [40, 330])
    d.kv_line([("Website", v["website"])], [40])
    d.kv_line([("Contact Name", v["contact"]), ("Contact Phone", v["contact_phone"])], [40, 330])
    d.kv_line([("Contact Email", v["contact_email"])], [40])
    d.heading("POLICY INFORMATION")
    d.kv_line([("Proposed Effective Date", v["eff"]), ("Proposed Expiration Date", v["exp"])], [40, 330])
    d.kv_line([("Transaction", v["transaction"]), ("Prior Carrier", v["prior_carrier"])], [40, 330])
    d.new_page(landscape_page=True)
    d.text(40, "PROPERTY SECTION", size=12, bold=True)
    d.heading("PREMISES / BUILDING INFORMATION")
    d.table(PREMISES_COLS, _premises_rows(w, g, rng), size=7.5, x0=36)
    if len(g.locations) > 3:
        d.text(40, f"Additional locations: see attached Statement of Values ({len(g.locations)} locations total).")
    d.kv_line([("Number of Locations", v["nlocs"]), ("Total Insured Value (All Locations)", v["tiv"])], [40, 330])
    d.heading("COVERAGES")
    d.kv_line([("Building Limit", v["bldg_limit"]), ("Business Personal Property Limit", v["bpp_limit"])], [40, 330])
    d.kv_line([("Business Income Limit", v["bi_limit"]), ("Deductible", v["deductible"])], [40, 330])
    d.kv_line([("Valuation", v["valuation"]), ("Coinsurance", v["coins"])], [40, 330])
    d.kv_line([("Causes of Loss Form", v["form"])], [40])
    d.kv_line([("Flood Coverage Requested", v["flood"]), ("Earthquake Coverage Requested", v["eq"])], [40, 330])
    d.new_page()
    _sign_page(d, v, g)
    return d.finish()


def render_application_b(w: World, g: Submission, rng: random.Random) -> bytes:
    v = _app_values(w, g, rng)
    d = PdfDoc()
    d.text(40, "APPLICATION FOR COMMERCIAL PROPERTY INSURANCE", size=13, bold=True)
    d.text(40, "Agency-generated form (synthetic, ACORD-like fields)", size=8)
    d.heading("APPLICANT")
    d.kv_line([("Applicant Name", v["named_insured"]), ("DBA / Trade Name", v["dba"])], [40, 330])
    d.kv_line([("FEIN/Tax ID", v["fein"]), ("Business Entity", v["entity"])], [40, 330])
    d.kv_line([("Street", v["street"]), ("City", v["city"])], [40, 330])
    d.kv_line([("State", v["state"]), ("ZIP", v["zip"])], [40, 330])
    d.kv_line([("Description of Operations", v["description"])], [40])
    d.kv_line([("NAICS", v["naics"]), ("Yrs in Bus.", v["years"])], [40, 330])
    d.kv_line([("Gross Revenue", v["revenue"]), ("Employees", v["employees"])], [40, 330])
    d.kv_line([("Web Address", v["website"]), ("Primary Contact", v["contact"])], [40, 330])
    d.kv_line([("Contact Phone", v["contact_phone"]), ("Contact E-mail", v["contact_email"])], [40, 330])
    d.heading("POLICY")
    d.kv_line([("Effective Date", v["eff"]), ("Expiration Date", v["exp"])], [40, 330])
    d.kv_line([("New / Renewal", v["transaction"]), ("Current Carrier", v["prior_carrier"])], [40, 330])
    d.heading("COVERAGE REQUESTED")
    d.kv_line([("Building", v["bldg_limit"]), ("Business Personal Property", v["bpp_limit"])], [40, 330])
    d.kv_line([("Business Income", v["bi_limit"]), ("Deductible", v["deductible"])], [40, 330])
    d.kv_line([("Valuation", v["valuation"]), ("Coinsurance", v["coins"])], [40, 330])
    d.kv_line([("Cause of Loss", v["form"]), ("Flood", v["flood"])], [40, 330])
    d.kv_line([("Earthquake", v["eq"]), ("Total TIV", v["tiv"])], [40, 330])
    d.heading("PRODUCER")
    d.kv_line([("Agency Name", v["agency"]), ("Producer Name", v["producer"])], [40, 330])
    d.kv_line([("Producer Phone", v["producer_phone"]), ("Producer E-mail", v["producer_email"])], [40, 330])
    d.new_page(landscape_page=True)
    d.heading("SCHEDULE OF LOCATIONS (FIRST THREE)")
    d.table(PREMISES_COLS, _premises_rows(w, g, rng), size=7.5, x0=36)
    d.kv_line([("Number of Locations", v["nlocs"])], [40])
    if len(g.locations) > 3:
        d.text(40, "Remaining locations per attached SOV.")
    d.new_page()
    _sign_page(d, v, g)
    return d.finish()


def render_application_c(w: World, g: Submission, rng: random.Random) -> bytes:
    """Held-out layout (test split only): boxed fields, label above value, different wording."""
    v = _app_values(w, g, rng)
    m = w.messy_formats
    d = PdfDoc()
    d.text(40, "PROPERTY RISK QUESTIONNAIRE", size=14, bold=True)
    d.text(40, "Wholesale broker intake form (synthetic)", size=8, dy=20)
    d.boxes([("Legal entity name", v["named_insured"]), ("Trade name", v["dba"])], [300, 232])
    d.boxes([("Tax ID number", v["fein"]), ("Organization", v["entity"]), ("Industry code", v["naics"])], [180, 176, 176])
    d.boxes([("Street address", v["street"]), ("City", v["city"]), ("State", v["state"]), ("Postal code", v["zip"])],
            [220, 150, 60, 102])
    d.boxes([("Nature of operations", v["description"])], [532])
    d.boxes([("Years in operation", v["years"]), ("Annual sales", v["revenue"]), ("Full-time staff", v["employees"])],
            [180, 176, 176])
    d.boxes([("Web", v["website"]), ("Inspection contact", v["contact"])], [266, 266])
    d.boxes([("Contact tel", v["contact_phone"]), ("Contact email", v["contact_email"])], [266, 266])
    d.boxes([("Inception", v["eff"]), ("Expiry", v["exp"]), ("Business type", v["transaction"])], [180, 176, 176])
    d.boxes([("Incumbent insurer", v["prior_carrier"]), ("Total values", v["tiv"]), ("Locations", v["nlocs"])],
            [220, 156, 156])
    d.boxes([("Bldg limit", v["bldg_limit"]), ("Contents limit", v["bpp_limit"]), ("Time element", v["bi_limit"])],
            [180, 176, 176])
    d.boxes([("AOP deductible", v["deductible"]), ("Valuation basis", v["valuation"]), ("Coins.", v["coins"])],
            [180, 176, 176])
    d.boxes([("Perils", v["form"]), ("Flood?", v["flood"]), ("Quake?", v["eq"])], [180, 176, 176])
    d.boxes([("Submitting broker", v["producer"]), ("Brokerage", v["agency"])], [220, 312])
    d.boxes([("Broker tel", v["producer_phone"]), ("Broker email", v["producer_email"])], [220, 312])
    d.new_page()
    d.text(40, "BUILDING DETAILS", size=12, bold=True, dy=18)
    for loc in g.locations[:3]:
        yb = w.app_year_built.get(loc.loc_number, loc.year_built)
        d.ensure(120)
        d.text(40, f"Building {loc.loc_number}", size=10, bold=True)
        d.boxes([("Address", _addr_line(loc.street, loc.city, loc.state, loc.zip))], [532])
        d.boxes([("Const. class", F.construction(loc.construction, rng, m)), ("Built", F.intval(yb)),
                 ("Floors", F.intval(loc.stories)), ("Area (sf)", F.intcomma(loc.square_feet))], [160, 124, 124, 124])
        d.boxes([("Sprinklers", F.yes_no(loc.sprinklered, rng, m)), ("Fire prot. class", F.intval(loc.protection_class)),
                 ("Roof covering", F.ROOF_LABEL.get(loc.roof_type or "", "")), ("Roof updated", F.intval(loc.roof_year)),
                 ("Burglar alarm", F.ALARM_LABEL.get(loc.alarm_type or "", ""))], [100, 108, 124, 100, 100])
    if len(g.locations) > 3:
        d.text(40, "Other buildings: refer to the statement of values.")
    d.new_page()
    _sign_page(d, v, g)
    return d.finish()


# --------------------------------------------------------------------------- loss run

LOSS_COLS = [Column("Claim No.", 100), Column("Date of Loss", 70), Column("Cause of Loss", 110),
             Column("Status", 50), Column("Paid", 80, "right"), Column("Reserve", 75, "right"),
             Column("Total Incurred", 85, "right")]


def render_loss_run(w: World, g: Submission, carrier: str, years: list[int], rng: random.Random) -> bytes:
    d = PdfDoc(font=rng.choice(["Helvetica", "Helvetica", "Courier"]))
    eff = g.policy.effective_date
    valuation = (g.policy.submission_received_date or eff) - timedelta(days=rng.randint(5, 30))
    d.text(40, carrier.upper(), size=13, bold=True)
    d.text(40, rng.choice(["LOSS RUN REPORT", "CLAIMS HISTORY - LOSS EXPERIENCE", "Loss History Report"]), size=11, bold=True)
    d.text(40, f"Insured: {w.name_variants['loss_run']}")
    d.text(40, f"Policy Number: CPP-{rng.randint(100000, 999999)}    Line: Commercial Property")
    d.text(40, f"Valuation Date: {valuation.strftime('%m/%d/%Y')}", dy=20)
    split_hint = rng.random() < 0.5
    for y in years:
        ly = next(l for l in g.loss_years if l.policy_year == y)
        start = eff.replace(year=y)
        end = eff.replace(year=y + 1)
        claims = [c for c in g.claims if c.policy_year == y]
        d.ensure(60)
        d.text(40, f"Policy Period: {start.strftime('%m/%d/%Y')} - {end.strftime('%m/%d/%Y')}    "
                   f"Earned Premium: {F.money_cents(ly.earned_premium)}", size=9, bold=True)
        if not claims:
            d.text(52, "No losses reported for this period.", dy=18)
            continue
        rows = [[c.claim_number, c.date_of_loss.strftime("%m/%d/%Y"), c.cause or "", (c.status or "").title(),
                 F.money_cents(c.paid), F.money_cents(c.reserve), F.money_cents(c.incurred)] for c in claims]
        d.table(LOSS_COLS, rows, size=8, row_h=14, repeat_header=not split_hint, x0=40)
        d.text(52, f"Period Total: {len(claims)} claim(s)    Total Incurred: {F.money_cents(ly.total_incurred)}",
               size=8.5, bold=True, dy=20)
    d.ensure(40)
    d.text(40, "This report reflects claims information as of the valuation date shown.", size=7)
    return d.finish()


# --------------------------------------------------------------------------- SOV

SOV_HEADERS = {
    "loc": ["Loc #", "Location No.", "Loc"],
    "street": ["Street Address", "Address", "Location Address"],
    "city": ["City"],
    "state": ["State", "ST"],
    "zip": ["Zip", "Zip Code", "Postal"],
    "occupancy": ["Occupancy", "Occupancy / Use", "Building Use"],
    "construction": ["Construction", "Const Type", "ISO Construction"],
    "year_built": ["Year Built", "Yr Built", "Yr Blt"],
    "stories": ["Stories", "# Stories", "No. of Floors"],
    "square_feet": ["Sq Ft", "Square Footage", "Area (SF)"],
    "sprinklered": ["Sprinklered", "Sprinklered (Y/N)", "Sprk"],
    "sprinkler_pct": ["Sprinkler %", "% Sprinklered", "Sprk %"],
    "protection_class": ["Protection Class", "PPC", "Prot Class"],
    "roof_year": ["Roof Year", "Roof Updated", "Yr Roof"],
    "building_value": ["Building Value", "Building Limit", "Bldg Value"],
    "contents_value": ["Contents Value", "BPP", "Contents"],
    "bi_value": ["BI Value", "Business Income", "BI/EE"],
    "tiv": ["TIV", "Total Insured Value", "Total Value"],
}
SOV_ORDER = ["loc", "street", "city", "state", "zip", "occupancy", "construction", "year_built", "stories",
             "square_feet", "sprinklered", "sprinkler_pct", "protection_class", "roof_year",
             "building_value", "contents_value", "bi_value", "tiv"]


def _sov_value(key: str, loc, w: World, rng: random.Random, as_number: bool):
    path = f"locations[{loc.loc_number}].{key}"
    if path in w.omitted:
        return None
    v = loc.loc_number if key == "loc" else getattr(loc, key)
    if v is None:
        return None
    m = w.messy_formats
    if key == "construction":
        return F.construction(v, rng, m)
    if key == "sprinklered":
        return F.yes_no(v, rng, m)
    if key == "sprinkler_pct":
        return v if as_number else f"{v}%"
    if key in ("building_value", "contents_value", "bi_value", "tiv"):
        if as_number and not (m and rng.random() < 0.15):
            return float(v)
        return F.money(v, rng, m)
    if key == "square_feet":
        return v if as_number else F.intcomma(v)
    return v


def _sov_layout(rng: random.Random) -> tuple[list[str], dict[str, str]]:
    order = list(SOV_ORDER)
    if rng.random() < 0.3:  # some brokers put values first
        values = order[-4:]
        order = order[:6] + values + order[6:-4]
    headers = {k: rng.choice(v) for k, v in SOV_HEADERS.items()}
    return order, headers


def render_sov_xlsx(w: World, g: Submission, rng: random.Random) -> bytes:
    order, headers = _sov_layout(rng)
    wb = Workbook()
    ws = wb.active
    ws.title = rng.choice(["SOV", "Schedule", "Sheet1", "Locations"])
    r = 1
    if rng.random() < 0.7:
        ws.cell(r, 1, f"Statement of Values - {w.name_variants['sov']}").font = Font(bold=True, size=13)
        r += 1
        ws.cell(r, 1, f"As of {g.policy.submission_received_date.strftime('%m/%d/%Y')}")
        r += 2
    extra = rng.random() < 0.3
    hdrs = [headers[k] for k in order] + (["Notes"] if extra else [])
    for j, h in enumerate(hdrs, start=1):
        ws.cell(r, j, h).font = Font(bold=True)
    r += 1
    for loc in g.locations:
        for j, k in enumerate(order, start=1):
            ws.cell(r, j, _sov_value(k, loc, w, rng, as_number=True))
        if extra and rng.random() < 0.3:
            ws.cell(r, len(order) + 1, rng.choice(["Updated 2024", "Vacant 2nd floor", "New roof", "Tenant occupied"]))
        r += 1
    if rng.random() < 0.6:
        ws.cell(r, 1, "TOTAL").font = Font(bold=True)
        j_tiv = order.index("tiv") + 1
        ws.cell(r, j_tiv, float(g.totals.total_tiv or 0)).font = Font(bold=True)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_sov_pdf(w: World, g: Submission, rng: random.Random) -> bytes:
    order, headers = _sov_layout(rng)
    widths = {"loc": 24, "street": 96, "city": 58, "state": 20, "zip": 30, "occupancy": 58, "construction": 66,
              "year_built": 30, "stories": 32, "square_feet": 40, "sprinklered": 34, "sprinkler_pct": 34,
              "protection_class": 36, "roof_year": 38, "building_value": 56, "contents_value": 50,
              "bi_value": 46, "tiv": 56}
    right = {"square_feet", "building_value", "contents_value", "bi_value", "tiv"}
    total_w = sum(widths.values())
    scale = min(1.0, 740 / total_w)
    cols = [Column(_short(headers[k]), widths[k] * scale, "right" if k in right else "left") for k in order]
    d = PdfDoc(landscape_first=True)
    d.text(30, f"STATEMENT OF VALUES - {w.name_variants['sov'].upper()}", size=11, bold=True)
    d.text(30, f"Prepared by {g.broker.agency_name}", size=8, dy=16)
    rows = []
    for loc in g.locations:
        row = []
        for k in order:
            val = _sov_value(k, loc, w, rng, as_number=False)
            row.append("" if val is None else str(val))
        rows.append(row)
    force_split = len(rows) >= 5 and rng.random() < 0.6
    d.table(cols, rows, size=6.5, row_h=13, repeat_header=rng.random() < 0.5, landscape_pages=True,
            max_rows_first_page=(len(rows) // 2) if force_split else None, x0=24)
    d.text(30, f"Total Insured Value: {F.money(g.totals.total_tiv, rng, False)}", size=8, bold=True)
    return d.finish()


def _short(h: str) -> str:
    return {"Street Address": "Street Address", "Location Address": "Address", "Total Insured Value": "TIV",
            "Sprinklered (Y/N)": "Sprk Y/N", "Square Footage": "Sq Ft", "ISO Construction": "Construction",
            "Occupancy / Use": "Occupancy", "Business Income": "BI", "No. of Floors": "Floors",
            "% Sprinklered": "% Sprk", "Roof Updated": "Roof Upd", "Protection Class": "Prot. Class",
            "Building Value": "Bldg Value", "Building Limit": "Bldg Limit", "Contents Value": "Contents"}.get(h, h)
